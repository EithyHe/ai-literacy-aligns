"""Report generation service (PRD FR-4.4.1-4.4.5).

Generates standardized semantic analysis reports in Markdown, Word, and PDF.
"""

from __future__ import annotations

import io
import logging
from datetime import datetime
from typing import Any

from ssn.config import EMBEDDING_DIM, STREAMLIT_TITLE

logger = logging.getLogger(__name__)


def _section_header(title: str, level: int = 2) -> str:
    return f"{'#' * level} {title}\n\n"


def _format_table(headers: list[str], rows: list[list[Any]]) -> str:
    """Format a markdown table."""
    lines = ["| " + " | ".join(str(h) for h in headers) + " |"]
    lines.append("| " + " | ".join("---" for _ in headers) + " |")
    for row in rows:
        lines.append("| " + " | ".join(str(c) for c in row) + " |")
    return "\n".join(lines) + "\n\n"


def generate_method_section(
    embedding_model: str = "dwulff/mpnet-personality",
    corpus_size: int = 0,
    similarity_metric: str = "cosine",
) -> str:
    """Generate the standardised method description paragraph."""
    return (
        _section_header("Method: Semantic Analysis")
        + f"Semantic similarity analysis was conducted using the {STREAMLIT_TITLE} platform. "
        f"Item-level embeddings were generated using the **{embedding_model}** model "
        f"(embedding dimensionality = {EMBEDDING_DIM}). "
        f"The reference corpus contained **{corpus_size:,}** items from publicly available "
        f"psychological scales (IPIP repository). "
        f"Pairwise similarity was computed using **{similarity_metric}** similarity. "
        f"Nearest-neighbour analysis identified the most semantically similar existing "
        f"scales/constructs for each input scale, and item-level redundancy was assessed "
        f"using a three-tier risk classification (high: ≥ 0.85, medium: 0.65–0.85, low: < 0.65).\n\n"
    )


def generate_neighbors_section(neighbors: list[dict]) -> str:
    """Generate the neighbor scale comparison table."""
    if not neighbors:
        return _section_header("Neighbor Scale Comparison") + "_No results available._\n\n"
    headers = ["Rank", "Construct / Scale", "Framework", "Similarity"]
    rows = []
    for i, nb in enumerate(neighbors, 1):
        rows.append([
            i,
            nb.get("name", nb.get("construct_id", "")),
            nb.get("framework", ""),
            f"{nb.get('similarity', 0):.4f}",
        ])
    return _section_header("Neighbor Scale Comparison") + _format_table(headers, rows)


def generate_diagnosis_section(diagnosis: dict) -> str:
    """Generate the per-item diagnosis table."""
    if not diagnosis or "items" not in diagnosis:
        return _section_header("Item-Level Diagnosis") + "_No results available._\n\n"

    summary = diagnosis.get("summary", {})
    text = _section_header("Item-Level Diagnosis")
    text += (
        f"**Summary:** {summary.get('total', 0)} items analysed — "
        f"High risk: {summary.get('high_risk', 0)}, "
        f"Medium risk: {summary.get('medium_risk', 0)}, "
        f"Low risk: {summary.get('low_risk', 0)}. "
        f"Overall risk score: {summary.get('overall_risk_score', 0):.2f}\n\n"
    )

    headers = ["#", "Item Text", "Risk", "Most Similar Item", "Similarity"]
    rows = []
    for i, item in enumerate(diagnosis["items"], 1):
        top = item.get("top_matches", [{}])[0] if item.get("top_matches") else {}
        rows.append([
            i,
            item.get("item_text", "")[:60] + ("…" if len(item.get("item_text", "")) > 60 else ""),
            item.get("risk_level", ""),
            top.get("text", "")[:50] + ("…" if len(top.get("text", "")) > 50 else ""),
            f"{top.get('similarity', 0):.4f}" if top else "",
        ])
    text += _format_table(headers, rows)
    return text


def generate_density_section(density_results: dict) -> str:
    """Generate the density analysis section."""
    if not density_results:
        return _section_header("Density Analysis") + "_No results available._\n\n"
    text = _section_header("Density Analysis")
    if "density_percentile" in density_results:
        text += f"**Density percentile:** {density_results['density_percentile']:.1f}%\n\n"
    if "novelty_assessment" in density_results:
        text += f"**Novelty assessment:** {density_results['novelty_assessment']}\n\n"
    return text


def generate_dimensions_section(clusters: list[dict]) -> str:
    """Generate internal dimensions exploration section."""
    if not clusters:
        return _section_header("Internal Dimension Exploration") + "_No results available._\n\n"
    text = _section_header("Internal Dimension Exploration")
    headers = ["Cluster", "Items", "Theme"]
    rows = []
    for cl in clusters:
        rows.append([
            cl.get("cluster_id", ""),
            cl.get("item_count", 0),
            cl.get("theme", ""),
        ])
    text += _format_table(headers, rows)
    return text


def generate_corpus_section(stats: dict) -> str:
    """Generate corpus information section."""
    text = _section_header("Corpus Information")
    text += (
        f"- **Frameworks:** {stats.get('frameworks', 0)}\n"
        f"- **Domains:** {stats.get('domains', 0)}\n"
        f"- **Constructs:** {stats.get('constructs', 0)}\n"
        f"- **Items:** {stats.get('items', 0)}\n"
        f"- **Instruments:** {stats.get('instruments', 0)}\n\n"
    )
    return text


def generate_citation_section() -> str:
    """Generate the citation recommendation."""
    text = _section_header("Citation")
    text += (
        "Please cite this tool as:\n\n"
        "> Semantic Scale Network (LLM-based). "
        "Available at: https://github.com/XinyiHeFengJi/ai-literacy-aligns\n\n"
        "Original method:\n\n"
        "> Rosenbusch, H., Wanders, F., & Pit, I. L. (2020). The Semantic Scale Network. "
        "*Psychological Methods, 25*(3), 380–392.\n\n"
    )
    return text


def generate_full_report(
    corpus_stats: dict,
    neighbors: list[dict] | None = None,
    diagnosis: dict | None = None,
    density_results: dict | None = None,
    dimension_clusters: list[dict] | None = None,
    embedding_model: str = "dwulff/mpnet-personality",
    similarity_metric: str = "cosine",
) -> str:
    """Generate the full markdown report combining all sections."""
    report = f"# Semantic Analysis Report\n\n"
    report += f"**Generated:** {datetime.now().strftime('%Y-%m-%d %H:%M')}\n\n"
    report += "---\n\n"

    report += generate_method_section(
        embedding_model=embedding_model,
        corpus_size=corpus_stats.get("items", 0),
        similarity_metric=similarity_metric,
    )
    report += generate_neighbors_section(neighbors or [])
    report += generate_diagnosis_section(diagnosis or {})
    report += generate_density_section(density_results or {})
    report += generate_dimensions_section(dimension_clusters or [])
    report += generate_corpus_section(corpus_stats)
    report += generate_citation_section()

    report += "---\n\n"
    report += _section_header("Reproducibility Information", level=2)
    report += (
        f"- Embedding model: {embedding_model}\n"
        f"- Embedding dimensionality: {EMBEDDING_DIM}\n"
        f"- Corpus version: {datetime.now().strftime('%Y-%m-%d')} "
        f"({corpus_stats.get('items', 0):,} items)\n"
        f"- Similarity metric: {similarity_metric}\n"
        f"- Redundancy thresholds: high ≥ 0.85, medium ≥ 0.65\n"
    )

    return report


def export_to_word(report_md: str) -> bytes:
    """Convert markdown report to Word (.docx) bytes."""
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        raise ImportError("python-docx is required for Word export. Install with: pip install python-docx")

    doc = Document()
    for line in report_md.split("\n"):
        stripped = line.strip()
        if stripped.startswith("# ") and not stripped.startswith("## "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("| "):
            doc.add_paragraph(stripped, style="List Bullet")
        elif stripped.startswith("- "):
            doc.add_paragraph(stripped[2:], style="List Bullet")
        elif stripped.startswith("> "):
            p = doc.add_paragraph(stripped[2:])
            p.style = "Intense Quote" if "Intense Quote" in [s.name for s in doc.styles] else "Quote"
        elif stripped == "---":
            doc.add_page_break()
        elif stripped:
            doc.add_paragraph(stripped)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def export_to_pdf(report_md: str) -> bytes:
    """Convert markdown report to PDF bytes via WeasyPrint.

    Requires weasyprint. Falls back to HTML intermediate (WeasyPrint renders HTML).
    """
    try:
        from weasyprint import HTML, CSS
    except ImportError:
        raise ImportError(
            "weasyprint is required for PDF export. Install with: pip install weasyprint"
        )

    # Simple markdown-to-HTML: minimal conversion for WeasyPrint
    html_content = _markdown_to_html(report_md)
    html_doc = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="utf-8">
        <style>
            body {{ font-family: Georgia, serif; margin: 2em; line-height: 1.5; }}
            h1 {{ font-size: 1.5em; border-bottom: 1px solid #ccc; }}
            h2 {{ font-size: 1.2em; margin-top: 1.2em; }}
            table {{ border-collapse: collapse; margin: 1em 0; }}
            th, td {{ border: 1px solid #ddd; padding: 6px 10px; text-align: left; }}
            th {{ background: #f5f5f5; }}
        </style>
    </head>
    <body>
    {html_content}
    </body>
    </html>
    """
    buf = io.BytesIO()
    HTML(string=html_doc).write_pdf(buf)
    buf.seek(0)
    return buf.read()


def _markdown_to_html(md: str) -> str:
    """Minimal markdown to HTML for report body (headings, lists, paragraphs, tables)."""
    lines = md.replace("\r\n", "\n").split("\n")
    out: list[str] = []
    in_table = False
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("# "):
            out.append(f"<h1>{_escape(stripped[2:])}</h1>")
        elif stripped.startswith("## "):
            out.append(f"<h2>{_escape(stripped[3:])}</h2>")
        elif stripped.startswith("### "):
            out.append(f"<h3>{_escape(stripped[4:])}</h3>")
        elif stripped.startswith("| "):
            if not in_table:
                in_table = True
                out.append("<table>")
            cells = [c.strip() for c in stripped.split("|")[1:-1]]
            next_stripped = lines[i + 1].strip() if i + 1 < len(lines) else ""
            is_sep = "---" in next_stripped and next_stripped.replace("-", "").replace("|", "").strip() == ""
            if is_sep:
                out.append("<tr>" + "".join(f"<th>{_escape(c)}</th>" for c in cells) + "</tr>")
                i += 2
                continue
            out.append("<tr>" + "".join(f"<td>{_escape(c)}</td>" for c in cells) + "</tr>")
        elif stripped.startswith("- ") or stripped.startswith("* "):
            out.append(f"<li>{_escape(stripped[2:])}</li>")
        elif stripped.startswith("> "):
            out.append(f"<blockquote>{_escape(stripped[2:])}</blockquote>")
        elif stripped == "---":
            out.append("<hr>")
        elif in_table:
            in_table = False
            out.append("</table>")
            if stripped:
                out.append(f"<p>{_escape(stripped)}</p>")
        elif stripped:
            out.append(f"<p>{_escape(stripped)}</p>")
        i += 1
    if in_table:
        out.append("</table>")
    return "\n".join(out)


def _escape(s: str) -> str:
    return (
        s.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )
